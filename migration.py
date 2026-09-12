import openpyxl, csv, os, json, re

XLSX = "মাসিক হিসাব-2020-2025-New.xlsx"
OUT = "migration_output"
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
months = list(range(1,13))
annual_sheets = [2021,2022,2023,2024,2025]

# Collect members by serial/code. Names in the workbook may be Bijoy-encoded.
members = {}
for y in annual_sheets:
    ws = wb[str(y)]
    for r in range(4, ws.max_row+1):
        code, name = ws.cell(r,1).value, ws.cell(r,2).value
        if code is not None and name:
            members[str(int(code))] = str(name).strip()

with open(os.path.join(OUT,"members.csv"),"w",newline="",encoding="utf-8-sig") as f:
    w=csv.writer(f); w.writerow(["member_code","full_name"])
    for code,name in sorted(members.items(), key=lambda x:int(x[0])):
        w.writerow([code,name])

with open(os.path.join(OUT,"member_payments.csv"),"w",newline="",encoding="utf-8-sig") as f:
    w=csv.writer(f); w.writerow(["member_code","payment_year","payment_month","amount"])
    for y in annual_sheets:
        ws=wb[str(y)]
        for r in range(4,ws.max_row+1):
            code=ws.cell(r,1).value
            if code is None: continue
            code=str(int(code))
            for m in months:
                v=ws.cell(r,2+m).value
                if isinstance(v,(int,float)) and v>0:
                    w.writerow([code,y,m,float(v)])

print("Created:", os.listdir(OUT))
print("Note: verify names before import. The workbook appears to use Bijoy/legacy encoded Bengali text.")
